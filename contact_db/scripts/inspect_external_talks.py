
#!/usr/bin/env python3
from __future__ import annotations

import sys
from pathlib import Path

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))


def main():
    excel_path = project_root / "external_data" / "KolClaw达人标签&沟通话术.xlsx"

    if not excel_path.exists():
        print(f"文件不存在: {excel_path}")
        return

    print(f"文件位置: {excel_path}")
    print(f"文件大小: {excel_path.stat().st_size / 1024:.2f} KB")
    print()

    try:
        import openpyxl

        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True, keep_vba=False)
        print(f"Sheet 名称: {wb.sheetnames}")
        print()

        for sheet_name in wb.sheetnames:
            print("=" * 60)
            print(f"Sheet: {sheet_name}")
            print("=" * 60)

            ws = wb[sheet_name]
            rows_read = 0
            for i, row in enumerate(ws.iter_rows(values_only=True), 1):
                if i <= 10:
                    print(f"Row {i}: {row}")
                rows_read += 1
            print()
            print(f"总行数: {rows_read}")
            print()

    except ImportError:
        print("需要安装 openpyxl: pip install openpyxl")
        print("或者使用 pandas: pip install pandas openpyxl")
        print()
        print("尝试用 pandas 读取...")

        try:
            import pandas as pd

            xl = pd.ExcelFile(excel_path)
            print(f"Sheet 名称: {xl.sheet_names}")
            print()

            for sheet_name in xl.sheet_names:
                print("=" * 60)
                print(f"Sheet: {sheet_name}")
                print("=" * 60)
                df = pd.read_excel(excel_path, sheet_name=sheet_name)
                print(df.head(10).to_string())
                print()
                print(f"形状: {df.shape}")
                print(f"列名: {list(df.columns)}")
                print()

        except ImportError:
            print("也没有 pandas，建议安装: pip install pandas openpyxl")


if __name__ == "__main__":
    main()

